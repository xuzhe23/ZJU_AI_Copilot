# -*- coding: utf-8 -*-
"""文档解析：按扩展名抽取正文并切片分块；图片走 OCR；PDF/Office 为可选依赖。"""
from __future__ import annotations

import csv
import json
from pathlib import Path

SUPPORTED_TEXT = {".txt", ".md", ".markdown", ".log", ".py", ".yaml", ".yml", ".json", ".jsonl", ".csv"}
SUPPORTED_IMAGE = {".png", ".jpg", ".jpeg", ".bmp", ".webp", ".tif", ".tiff"}
SUPPORTED_OFFICE = {".pdf", ".docx", ".xlsx", ".xls", ".doc"}
SUPPORTED_SUFFIXES = SUPPORTED_TEXT | SUPPORTED_IMAGE | SUPPORTED_OFFICE


def parse_file(path: Path, ocr=None) -> str:
    """按扩展名解析文件，返回文本。"""
    ext = path.suffix.lower()
    try:
        if ext in SUPPORTED_IMAGE:
            return parse_image(path, ocr)
        if ext == ".csv":
            return parse_csv(path)
        if ext == ".json":
            return parse_json(path)
        if ext == ".jsonl":
            return path.read_text(encoding="utf-8", errors="ignore")
        if ext == ".pdf":
            return parse_pdf(path)
        if ext == ".docx":
            return parse_docx(path)
        if ext in {".xlsx", ".xls"}:
            return parse_xlsx(path)
        # 默认按文本读取
        return path.read_text(encoding="utf-8", errors="ignore")
    except Exception as e:  # noqa: BLE001
        return f"[解析失败 {path.name}: {e}]"


def parse_image(path: Path, ocr=None) -> str:
    if ocr is None:
        return f"[图片 {path.name} 未配置 OCR 引擎]"
    try:
        return ocr.recognize(str(path))
    except Exception as e:  # noqa: BLE001
        return f"[OCR 失败 {path.name}: {e}]"


def parse_csv(path: Path) -> str:
    rows = []
    with open(path, newline="", encoding="utf-8", errors="ignore") as f:
        for row in csv.reader(f):
            rows.append(" | ".join(cell.strip() for cell in row))
    return "\n".join(rows)


def parse_json(path: Path) -> str:
    data = json.loads(path.read_text(encoding="utf-8", errors="ignore"))
    return json.dumps(data, ensure_ascii=False, indent=2)


def parse_pdf(path: Path) -> str:
    try:
        import pypdf  # 可选
        reader = pypdf.PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except ImportError:
        try:
            import fitz  # PyMuPDF 可选
            doc = fitz.open(str(path))
            return "\n".join(page.get_text() for page in doc)
        except ImportError:
            return f"[PDF {path.name}：缺少 pypdf 或 PyMuPDF，请 pip install pypdf]"


def parse_docx(path: Path) -> str:
    try:
        import docx  # 可选
        d = docx.Document(str(path))
        return "\n".join(p.text for p in d.paragraphs)
    except ImportError:
        return f"[DOCX {path.name}：缺少 python-docx，请 pip install python-docx]"


def parse_xlsx(path: Path) -> str:
    try:
        import openpyxl  # 可选
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                out.append(" | ".join("" if c is None else str(c) for c in row))
        return "\n".join(out)
    except ImportError:
        return f"[Excel {path.name}：缺少 openpyxl，请 pip install openpyxl]"


def chunk_text(text: str, size: int = 500, overlap: int = 60) -> list:
    """按字符滑窗分块（对中文友好），size/overlap 可调。"""
    text = text.strip()
    if not text:
        return []
    chunks, start, n = [], 0, len(text)
    while start < n:
        end = min(start + size, n)
        chunks.append(text[start:end])
        if end >= n:
            break
        start = end - overlap
    return chunks
